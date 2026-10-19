from datetime import date

import pytest
from docx import Document

from bank_agent.config import Settings
from bank_agent.knowledge import KnowledgeStore
from bank_agent.schemas import DocumentMetadata


def metadata(policy_id):
    return DocumentMetadata(policy_id=policy_id, title="解析测试教学制度", version="1.0",
                            effective_date=date(2025, 1, 1), reviewed=True, reviewed_by="fixture-reviewer",
                            document_type="teaching")


def write_pdf(path, text):
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii") if text else b""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    body = b"%PDF-1.4\n"
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(body))
        body += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(body)
    body += b"xref\n0 6\n0000000000 65535 f \n"
    body += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    body += f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    path.write_bytes(body)


def test_word_preserves_order_table_headers_numbers_and_no_invented_pages(tmp_path):
    path = tmp_path / "matrix.docx"
    document = Document()
    document.add_paragraph("第五条 授权矩阵如下，口头批准不得替代书面批准。")
    table = document.add_table(rows=2, cols=3)
    for row, values in zip(table.rows, [("金额", "审批人", "例外"), ("500万元", "授权委员会", "不得口头批准")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.add_paragraph("第六条 申请记录应保存30日。")
    document.save(path)
    store = KnowledgeStore(Settings(_env_file=None, data_dir=tmp_path / "store"))
    clauses = store.ingest(path, metadata("word-fixture"))
    assert [clause.article for clause in clauses] == ["第五条", "第五条·表1", "第六条"]
    assert "金额 | 审批人 | 例外" in clauses[1].text
    assert "500万元 | 授权委员会 | 不得口头批准" in clauses[1].text
    assert clauses[1].evidence_id in clauses[0].related_ids
    assert all(clause.page is None for clause in clauses)


def test_text_pdf_preserves_page_number_and_negation(tmp_path):
    path = tmp_path / "rule.pdf"
    write_pdf(path, "The amount shall not exceed 1000 USD.")
    store = KnowledgeStore(Settings(_env_file=None, data_dir=tmp_path / "store"))
    clauses = store.ingest(path, metadata("pdf-fixture"))
    assert len(clauses) == 1
    assert clauses[0].text == "The amount shall not exceed 1000 USD."
    assert clauses[0].page == 1


def test_pdf_without_text_explicitly_requires_ocr(tmp_path):
    path = tmp_path / "blank.pdf"
    write_pdf(path, "")
    store = KnowledgeStore(Settings(_env_file=None, data_dir=tmp_path / "store"))
    with pytest.raises(ValueError, match="OCR"):
        store.ingest(path, metadata("blank-fixture"))
