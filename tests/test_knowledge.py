import hashlib
import json
import sqlite3
from datetime import date

import pytest

from bank_agent.config import Settings
from bank_agent.knowledge import KnowledgeStore
from bank_agent.schemas import DocumentMetadata, QueryRequest, UserContext


@pytest.fixture
def store(tmp_path):
    instance = KnowledgeStore(Settings(data_dir=tmp_path / "data", _env_file=None))
    instance.register_index_builder(lambda snapshot_id: None)
    return instance


def metadata(policy_id="p", version="1", effective="2025-01-01", **kwargs):
    return DocumentMetadata(policy_id=policy_id, title="测试制度", version=version,
                            effective_date=date.fromisoformat(effective), reviewed=True,
                            reviewed_by="test-reviewer", **kwargs)


def ingest(store, tmp_path, meta, text="第一条 必须双人复核。\n第二条 必须保存记录。"):
    path = tmp_path / f"{meta.policy_id}-{meta.version}.txt"
    path.write_text(text, encoding="utf-8")
    return store.ingest(path, meta)


def query(day="2026-01-01", **kwargs):
    return QueryRequest(question="办理要求", business_date=date.fromisoformat(day), **kwargs)


def test_review_gate_and_immutable_version(store, tmp_path):
    path = tmp_path / "draft.txt"
    path.write_text("第一条 草稿。", encoding="utf-8")
    unreviewed = metadata().model_copy(update={"reviewed": False, "reviewed_by": ""})
    with pytest.raises(ValueError, match="人工审核"):
        store.ingest(path, unreviewed)
    original = ingest(store, tmp_path, metadata())
    repeated = ingest(store, tmp_path, metadata())
    assert original == repeated
    with pytest.raises(ValueError, match="不能覆盖"):
        ingest(store, tmp_path, metadata(), "第一条 修改了金额限制。")
    with pytest.raises(ValueError, match="不能覆盖"):
        ingest(store, tmp_path, metadata(access_level="restricted"))


def test_legacy_original_without_provenance_fields_can_be_reimported_without_changing_version(store, tmp_path):
    meta = metadata()
    original = ingest(store, tmp_path, meta)
    legacy_metadata = meta.model_dump(mode="json")
    legacy_metadata.pop("content_kind")
    legacy_metadata.pop("source_effective_date")
    legacy_clauses = [clause.model_dump(mode="json") for clause in original]
    for clause in legacy_clauses:
        clause["metadata"].pop("content_kind")
        clause["metadata"].pop("source_effective_date")
    legacy_payload = {
        "blocks": [{"text": text, "page": None, "table": False} for text in [
            "第一条 必须双人复核。", "第二条 必须保存记录。",
        ]],
        "metadata": legacy_metadata,
    }
    legacy_digest = hashlib.sha256(json.dumps(
        legacy_payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"),
    ).encode()).hexdigest()
    with sqlite3.connect(store.db_path) as connection:
        connection.execute(
            "UPDATE documents SET digest=?, metadata_json=?, clauses_json=? WHERE policy_id=? AND version=?",
            (legacy_digest, json.dumps(legacy_metadata, ensure_ascii=False),
             json.dumps(legacy_clauses, ensure_ascii=False), meta.policy_id, meta.version),
        )
    repeated = ingest(store, tmp_path, meta)
    assert repeated == original
    assert all(clause.metadata.content_kind == "original" for clause in repeated)
    assert all(clause.metadata.source_effective_date is None for clause in repeated)
    for changed in [
        meta.model_copy(update={"content_kind": "summary"}),
        meta.model_copy(update={"source_effective_date": date(2024, 1, 1)}),
    ]:
        with pytest.raises(ValueError, match="不能覆盖"):
            ingest(store, tmp_path, changed)
    with sqlite3.connect(store.db_path) as connection:
        assert connection.execute("SELECT COUNT(*) FROM documents").fetchone()[0] == 1
        assert connection.execute("SELECT digest FROM documents").fetchone()[0] == legacy_digest


def test_snapshot_publication_failure_and_rollback(store, tmp_path):
    original = ingest(store, tmp_path, metadata())
    first = store.publish()
    ingest(store, tmp_path, metadata(version="2", effective="2026-01-01"), "第一条 新版要求三人复核。")
    assert store.clauses() == original  # Ingest is draft until publishing.
    assert store.clauses(first) == original

    def failed_builder(snapshot_id):
        assert len(store.clauses(snapshot_id)) > len(original)
        raise RuntimeError("embedding unavailable")

    store.register_index_builder(failed_builder)
    with pytest.raises(RuntimeError, match="embedding unavailable"):
        store.publish()
    assert store.active_snapshot() == first
    failed = next(item for item in store.snapshots() if item["status"] == "failed")
    failed_clause = store.clauses(failed["snapshot_id"])[0]
    assert store.get_clause(failed_clause.evidence_id, UserContext(user_id="u"), query(),
                            failed["snapshot_id"]) is None
    assert not store.allowed_clauses(UserContext(user_id="u"), query(), failed["snapshot_id"])
    with pytest.raises(ValueError, match="ready"):
        store.activate(failed["snapshot_id"])
    store.register_index_builder(lambda snapshot_id: None)
    second = store.publish()
    assert store.active_snapshot() == second
    assert len(store.clauses(second)) > len(store.clauses(first))
    store.activate(first)
    assert store.active_snapshot() == first
    assert store.clauses() == original


def test_date_version_before_acl_and_no_expired_version_resurrection(store, tmp_path):
    ingest(store, tmp_path, metadata())
    ingest(store, tmp_path, metadata(version="2", effective="2026-01-01", access_level="restricted",
                                   expiry_date=date(2027, 1, 1)))
    store.publish()
    employee = UserContext(user_id="employee")
    compliance = UserContext(user_id="compliance", role="compliance")
    assert {c.metadata.version for c in store.allowed_clauses(employee, query("2025-12-31"))} == {"1"}
    assert not store.allowed_clauses(employee, query("2026-01-01"))
    assert {c.metadata.version for c in store.allowed_clauses(compliance, query("2026-01-01"))} == {"2"}
    assert not store.allowed_clauses(compliance, query("2027-01-01"))
    assert not store.allowed_clauses(compliance, query("2024-12-31"))


def test_organization_scope_role_and_teaching_prefilter(store, tmp_path):
    for policy_id, options in [
        ("general", {}), ("other_branch", {"organizations": ["branch-b"]}),
        ("restricted", {"access_level": "restricted"}),
        ("role_only", {"allowed_roles": ["admin"]}),
        ("loan", {"scopes": ["loan"]}),
        ("classroom", {"document_type": "teaching"}),
    ]:
        ingest(store, tmp_path, metadata(policy_id=policy_id, **options))
    store.publish()
    employee = UserContext(user_id="u", organization="branch-a")
    visible = store.allowed_clauses(employee, query(scope="deposit"))
    assert {c.metadata.policy_id for c in visible} == {"general"}
    teaching = store.allowed_clauses(employee, query(scope="deposit", include_teaching=True))
    assert {c.metadata.policy_id for c in teaching} == {"general", "classroom"}
    hidden = next(c for c in store.clauses() if c.metadata.policy_id == "other_branch")
    assert store.get_clause(hidden.evidence_id, employee, query()) is None


def test_docx_tables_and_related_articles_survive(tmp_path, store):
    from docx import Document

    document = Document()
    document.add_paragraph("第一章 柜面业务")
    document.add_paragraph("第一条 金额按下表执行；异常情形按第二条处理。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "金额"
    table.cell(0, 1).text = "要求"
    table.cell(1, 0).text = "10万元"
    table.cell(1, 1).text = "双人复核"
    document.add_paragraph("第二条 异常交易必须暂停办理并升级审批。")
    path = tmp_path / "tables.docx"
    document.save(path)
    clauses = store.ingest(path, metadata())
    first = next(c for c in clauses if c.article == "第一条")
    table = next(c for c in clauses if "·表" in c.article)
    second = next(c for c in clauses if c.article == "第二条")
    assert "金额 | 要求\n10万元 | 双人复核" == table.text
    assert table.evidence_id in first.related_ids
    assert first.evidence_id in table.related_ids
    assert second.evidence_id in first.related_ids
    assert "第一章 柜面业务" in first.text


def test_long_markdown_table_keeps_headers_and_rows_together(tmp_path, store):
    path = tmp_path / "table.md"
    rows = ["| 金额 | 审批要求 |", "| --- | --- |"]
    rows.extend(f"| {number}万元 | 双人复核且主管审批 |" for number in range(1, 80))
    path.write_text("## 第一条 按下表执行。\n" + "\n".join(rows) + "\n## 第二条 保存记录。", encoding="utf-8")
    clauses = store.ingest(path, metadata())
    table = next(c for c in clauses if "·表" in c.article)
    assert len(table.text) > store.settings.chunk_size
    assert table.text == "\n".join(rows)
    assert "片段" not in table.article
    first = next(c for c in clauses if c.article == "第一条")
    assert table.evidence_id in first.related_ids


def test_audit_is_allowlisted_and_redacted(store):
    store.audit({"question": "手机号13812345678，身份证110101199001011234，邮箱user@example.com，api_key=secret",
                 "user_id": "alice", "status": "answered", "evidence_ids": ["p:1:abc"],
                 "evidence": "complete confidential original", "answer": "sensitive answer",
                 "llm_api_key": "secret", "trace_id": "trace"})
    with sqlite3.connect(store.db_path) as connection:
        payload = connection.execute("SELECT payload_json FROM audit_log").fetchone()[0]
    assert "13812345678" not in payload
    assert "110101199001011234" not in payload
    assert "user@example.com" not in payload
    assert "secret" not in payload
    assert "complete confidential original" not in payload
    assert "sensitive answer" not in payload
    record = json.loads(payload)
    assert record["user_id"].startswith("sha256:")
    assert record["evidence_ids"] == ["p:1:abc"]


def test_publish_without_registered_index_is_rejected(tmp_path):
    store = KnowledgeStore(Settings(data_dir=tmp_path / "data", _env_file=None))
    ingest(store, tmp_path, metadata())
    with pytest.raises(RuntimeError, match="索引构建器"):
        store.publish()
    assert store.active_snapshot() == ""


def test_same_policy_same_effective_date_versions_cannot_publish(store, tmp_path):
    ingest(store, tmp_path, metadata())
    first = store.publish()
    ingest(store, tmp_path, metadata(version="2"))
    with pytest.raises(ValueError, match="生效日期相同"):
        store.publish()
    assert store.active_snapshot() == first


def test_audit_preserves_safe_nested_traces_but_drops_payloads(store):
    store.audit({"business_date": "2026-10-11", "rule_checks": ["passed"], "nodes": ["retrieve"],
                 "supplement_rounds": 1, "model_id": "model-revision", "repair_count": 0,
                 "versions": ["p@1"], "subquestions_used": 2,
                 "model_calls": [{"step": "analyze", "succeeded": True, "duration_ms": 50,
                                  "prompt": "confidential full prompt", "api_key": "secret"}],
                 "usage": {"calls": 1, "input_tokens": 100, "output_tokens": 20, "api_key": "secret"},
                 "retrieval_queries": [{"query": "手机号13812345678怎么办", "evidence_ids": ["p:1:id"],
                                        "evidence": "confidential evidence"}]})
    with sqlite3.connect(store.db_path) as connection:
        payload = connection.execute("SELECT payload_json FROM audit_log").fetchone()[0]
    record = json.loads(payload)
    assert record["model_calls"] == [{"step": "analyze", "succeeded": True, "duration_ms": 50}]
    assert record["usage"] == {"calls": 1, "input_tokens": 100, "output_tokens": 20}
    assert record["rule_checks"] == ["passed"]
    assert record["nodes"] == ["retrieve"]
    assert record["retrieval_queries"][0]["evidence_ids"] == ["p:1:id"]
    assert "13812345678" not in payload
    assert "secret" not in payload
    assert "confidential" not in payload
